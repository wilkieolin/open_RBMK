#!/usr/bin/env python3
"""
Render sources/derived/core_layout.xlsx (a hand digitisation of Dollezhal &
Emelyanov 1980, Fig. 2.3, p. 25 of the PDF: initial loading of Leningrad NPP
Unit 2) as a vector grid sheet, for overlay on the scanned figure, and check it
against the counts the same book gives in its text.

    python tools/render_core_layout.py        # -> sources/derived/core_layout_overlay.svg
                                              #    sources/derived/core_layout_checks.txt

Workbook layout (as digitised): row 1 = column coordinates 10..67 (RBMK
numbering, x8/x9 skipped), column A = row coordinates 67..10; a cell's text is
its symbol code from the legend in columns AY-BF.  Only a BLACK fill means
"outside the core"; any other fill (or none) inside the grid is a channel, and
the label alone decides its type.

Overlay: every symbol is drawn in one ink colour on a transparent background,
in the figure's own geometry (same channel coordinates, same label placement:
even columns on top, odd below; odd rows left, even right), so it can be scaled
onto the scan in any editor and misreadings show as red over black.  Layers
(grid / outline / symbols / labels / legend) can be toggled in Inkscape.
"""
import collections
import math
import os
import sys

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(ROOT, "sources", "derived", "core_layout.xlsx")
SVG = os.path.join(ROOT, "sources", "derived", "core_layout_overlay.svg")
REPORT = os.path.join(ROOT, "sources", "derived", "core_layout_checks.txt")

INK = "#d40000"
S = 20.0                      # cell size, SVG user units
MARGIN = 60.0

# legend: code -> (legend number, meaning); drawn by draw_symbol()
LEGEND = {
    "H": (1, "стержень РР (manual control rod)"),
    "Q": (2, "стержень УСП (shortened absorber rod)"),
    "XO": (3, "стержень АР (automatic rod)"),
    "F": (4, "стержень АЗ (emergency rod)"),
    "O": (5, "стержень перекомпенсации АЗ (AZ over-compensation rod)"),
    "X": (6, "ДКЭ(в) (axial power detector)"),
    "T": (7, "пусковая камера деления (start-up fission chamber)"),
    "D": (8, "дополнительный поглотитель (additional absorber)"),
    "*": (9, "незагруженный канал (unloaded channel)"),
    "o": (10, "ТВС с ДКЭ(р) (fuel with radial detector)"),
    "": (11, "тепловыделяющая кассета (fuel assembly)"),
}
RODS = {"H", "Q", "XO", "F", "O"}
SPECIAL = RODS | {"X", "T"}
TECH = {"", "o", "D", "*"}
# Checked against the scan by hand (2026-09-26), after the first digitisation:
#   corrected: 52-15, 52-21 fuel -> H; 46-11, 46-15 D -> H; 30-43 F -> O;
#   17-15 is outside the
#   core; 46-27 held a stray space (fuel).
#   confirmed real asymmetries (NOT errors): 40-63 = F with 14-37 = D, and
#   30-63 = O with 14-47 = D.  The start-up chambers (T) sit at the periphery.
#   radial-detector fuel (o): the figure has 282 against the text's 270 (p. 57,
#   "предусматривается").  Settled 2026-09-26 as a difference between the figure
#   and the text: the overlay shows no misreadings that could account for it.  It
#   does not matter to the model (these are ordinary fuel assemblies to neutronics).
CONFIRMED = {"40-63", "14-37", "30-63", "14-47", "67-47", "50-10", "30-67", "10-30"}
# Dollezhal & Emelyanov 1980, p. 57 (PDF) and the parameter table
EXPECT = {"technological channels": 1693, "CPS rods": 179, "DKE(v) channels": 12,
          "start-up chambers (KD)": 4}
# not enforced: the figure shows 282 radial-detector assemblies, the text 270 (see above)


def load():
    wb = openpyxl.load_workbook(XLSX)
    ws = wb.active
    cols = [ws.cell(1, c).value for c in range(2, 50)]
    rows = [ws.cell(r, 1).value for r in range(2, 50)]
    assert len(cols) == 48 and len(rows) == 48, (len(cols), len(rows))
    grid = {}
    for j, r in enumerate(range(2, 50)):
        for i, c in enumerate(range(2, 50)):
            cell = ws.cell(r, c)
            f = cell.fill
            black = (f.fill_type == "solid" and f.fgColor.type == "theme"
                     and f.fgColor.theme == 1 and abs(f.fgColor.tint or 0) < 1e-6) or \
                    (f.fill_type == "solid" and f.fgColor.type == "rgb"
                     and str(f.fgColor.rgb).upper().endswith("000000"))
            raw = cell.value
            v = "" if raw is None else str(raw).strip()
            grid[(i, j)] = {"out": black, "code": v, "raw": raw}
    return cols, rows, grid


# ----------------------------------------------------------------- drawing
def draw_symbol(code, cx, cy):
    r = 0.36 * S
    sw = 1.2
    circ = f'<circle cx="{cx:.2f}" cy="{cy:.2f}" r="{r:.2f}" fill="none" stroke="{INK}" stroke-width="{sw}"/>'
    if code == "H":        # half moon: right half filled
        return circ + (f'<path d="M {cx:.2f} {cy - r:.2f} A {r:.2f} {r:.2f} 0 0 1 {cx:.2f} {cy + r:.2f} Z" '
                       f'fill="{INK}"/>')
    if code == "Q":        # quarter moon: only the right-hand quadrant (SE to NE) filled
        d = r * math.sqrt(0.5)
        return circ + (f'<path d="M {cx:.2f} {cy:.2f} L {cx + d:.2f} {cy - d:.2f} A {r:.2f} {r:.2f} 0 0 1 '
                       f'{cx + d:.2f} {cy + d:.2f} Z" fill="{INK}"/>')
    if code == "XO":       # open circle with a cross
        d = r * math.sqrt(0.5)
        return circ + (f'<path d="M {cx - d:.2f} {cy - d:.2f} L {cx + d:.2f} {cy + d:.2f} '
                       f'M {cx - d:.2f} {cy + d:.2f} L {cx + d:.2f} {cy - d:.2f}" '
                       f'stroke="{INK}" stroke-width="{sw}"/>')
    if code == "F":        # full moon
        return f'<circle cx="{cx:.2f}" cy="{cy:.2f}" r="{r:.2f}" fill="{INK}"/>'
    if code == "O":        # open circle
        return circ
    if code == "X":        # as the AR symbol, with the left and right quadrants filled (as in the figure)
        d = r * math.sqrt(0.5)
        return (f'<path d="M {cx:.2f} {cy:.2f} L {cx + d:.2f} {cy - d:.2f} A {r:.2f} {r:.2f} 0 0 1 '
                f'{cx + d:.2f} {cy + d:.2f} Z M {cx:.2f} {cy:.2f} L {cx - d:.2f} {cy + d:.2f} '
                f'A {r:.2f} {r:.2f} 0 0 1 {cx - d:.2f} {cy - d:.2f} Z" fill="{INK}"/>'
                + circ + f'<path d="M {cx - d:.2f} {cy - d:.2f} L {cx + d:.2f} {cy + d:.2f} '
                f'M {cx - d:.2f} {cy + d:.2f} L {cx + d:.2f} {cy - d:.2f}" '
                f'stroke="{INK}" stroke-width="{sw}"/>')
    if code == "T":        # triangle in circle (point down)
        t = 0.8 * r
        return circ + (f'<path d="M {cx - t:.2f} {cy - 0.5 * t:.2f} L {cx + t:.2f} {cy - 0.5 * t:.2f} '
                       f'L {cx:.2f} {cy + t:.2f} Z" fill="{INK}"/>')
    if code == "D":        # Cyrillic П
        w, h = 0.28 * S, 0.34 * S
        return (f'<path d="M {cx - w:.2f} {cy + h:.2f} L {cx - w:.2f} {cy - h:.2f} L {cx + w:.2f} {cy - h:.2f} '
                f'L {cx + w:.2f} {cy + h:.2f}" fill="none" stroke="{INK}" stroke-width="{1.6 * sw}"/>')
    if code == "*":        # unloaded channel: cross corner to corner
        h = 0.5 * S
        return (f'<path d="M {cx - h:.2f} {cy - h:.2f} L {cx + h:.2f} {cy + h:.2f} '
                f'M {cx - h:.2f} {cy + h:.2f} L {cx + h:.2f} {cy - h:.2f}" '
                f'stroke="{INK}" stroke-width="{sw}"/>')
    if code == "o":        # dot
        return f'<circle cx="{cx:.2f}" cy="{cy:.2f}" r="{0.12 * S:.2f}" fill="{INK}"/>'
    if code == "":
        return ""
    # anything not in the legend: flag it
    h = 0.45 * S
    return (f'<rect x="{cx - h:.2f}" y="{cy - h:.2f}" width="{2 * h:.2f}" height="{2 * h:.2f}" '
            f'fill="none" stroke="#0070ff" stroke-width="2"/>'
            f'<text x="{cx:.2f}" y="{cy + 4:.2f}" font-size="11" text-anchor="middle" fill="#0070ff">?</text>')


def layer(name, body):
    return (f'<g inkscape:groupmode="layer" inkscape:label="{name}" id="{name}">\n'
            + "\n".join(body) + "\n</g>")


def render(cols, rows, grid):
    W = 2 * MARGIN + 48 * S
    H = 2 * MARGIN + 48 * S + 140
    x0 = y0 = MARGIN
    X = lambda i: x0 + i * S
    Y = lambda j: y0 + j * S
    # grid
    g = [f'<path d="' + " ".join(f"M {X(i):.1f} {Y(0):.1f} L {X(i):.1f} {Y(48):.1f}" for i in range(49))
         + " " + " ".join(f"M {X(0):.1f} {Y(j):.1f} L {X(48):.1f} {Y(j):.1f}" for j in range(49))
         + f'" stroke="{INK}" stroke-width="0.4" stroke-opacity="0.6" fill="none"/>']
    # outside-core tint and the core boundary
    out = [f'<rect x="{X(i):.1f}" y="{Y(j):.1f}" width="{S}" height="{S}" fill="{INK}" fill-opacity="0.12"/>'
           for (i, j), c in grid.items() if c["out"]]
    edges = []
    for (i, j), c in grid.items():
        if c["out"]:
            continue
        for di, dj, seg in ((1, 0, (X(i + 1), Y(j), X(i + 1), Y(j + 1))),
                            (-1, 0, (X(i), Y(j), X(i), Y(j + 1))),
                            (0, 1, (X(i), Y(j + 1), X(i + 1), Y(j + 1))),
                            (0, -1, (X(i), Y(j), X(i + 1), Y(j)))):
            nb = grid.get((i + di, j + dj))
            if nb is None or nb["out"]:
                edges.append("M {:.1f} {:.1f} L {:.1f} {:.1f}".format(*seg))
    out.append(f'<path d="{" ".join(edges)}" stroke="{INK}" stroke-width="2.4" fill="none" '
               f'stroke-linecap="square"/>')
    # symbols
    sym = [draw_symbol(c["code"], X(i) + S / 2, Y(j) + S / 2)
           for (i, j), c in grid.items() if not c["out"] and c["code"] != ""]
    sym += [draw_symbol("?", X(i) + S / 2, Y(j) + S / 2)
            for (i, j), c in grid.items() if c["out"] and c["code"] != ""]
    # labels, placed as in the figure
    lab = []
    for i, cv in enumerate(cols):
        if cv % 2 == 0:
            lab.append(f'<text x="{X(i) + S / 2:.1f}" y="{Y(0) - 8:.1f}" font-size="11" text-anchor="middle" '
                       f'fill="{INK}">{cv}</text>')
        else:
            lab.append(f'<text x="{X(i) + S / 2:.1f}" y="{Y(48) + 16:.1f}" font-size="11" text-anchor="middle" '
                       f'fill="{INK}">{cv}</text>')
    for j, rv in enumerate(rows):
        if rv % 2 == 1:
            lab.append(f'<text x="{X(0) - 8:.1f}" y="{Y(j) + S / 2 + 4:.1f}" font-size="11" text-anchor="end" '
                       f'fill="{INK}">{rv}</text>')
        else:
            lab.append(f'<text x="{X(48) + 8:.1f}" y="{Y(j) + S / 2 + 4:.1f}" font-size="11" '
                       f'fill="{INK}">{rv}</text>')
    # legend
    leg, lx, ly = [], x0, Y(48) + 40
    for k, (code, (n, txt)) in enumerate(LEGEND.items()):
        cx = lx + (k % 3) * 330
        cy = ly + (k // 3) * 24
        leg.append(f'<rect x="{cx - S / 2:.1f}" y="{cy - S / 2:.1f}" width="{S}" height="{S}" fill="none" '
                   f'stroke="{INK}" stroke-width="0.4"/>' + draw_symbol(code, cx, cy)
                   + f'<text x="{cx + 16:.1f}" y="{cy + 4:.1f}" font-size="9" fill="{INK}">{n} {txt}</text>')
    head = (f'<?xml version="1.0" encoding="UTF-8"?>\n'
            f'<svg xmlns="http://www.w3.org/2000/svg" '
            f'xmlns:inkscape="http://www.inkscape.org/namespaces/inkscape" '
            f'width="{W:.0f}" height="{H:.0f}" viewBox="0 0 {W:.0f} {H:.0f}" '
            f'font-family="DejaVu Sans, Arial, sans-serif">\n'
            f'<title>Overlay: core_layout.xlsx (after Dollezhal &amp; Emelyanov 1980, Fig. 2.3)</title>\n')
    return head + "\n".join([layer("grid", g), layer("outline", out), layer("symbols", sym),
                             layer("labels", lab), layer("legend", leg)]) + "\n</svg>\n"


# ------------------------------------------------------------------ checks
def checks(cols, rows, grid):
    L = []
    lab = lambda i, j: f"{rows[j]}-{cols[i]:02d}"
    inside = {k: c for k, c in grid.items() if not c["out"]}
    cnt = collections.Counter(c["code"] for c in inside.values())
    L.append("COUNTS (label-based; only black fill = outside)")
    L.append(f"  positions inside the core outline: {len(inside)}"
             f"   (expected {sum(EXPECT.values())} = 1693 + 179 + 12 + 4)")
    for code, (n, txt) in LEGEND.items():
        L.append(f"  {n:>2} {code or '(blank)':>7}  {cnt.get(code, 0):5d}  {txt}")
    odd = {c for c in cnt if c not in LEGEND}
    for c in sorted(odd):
        where = [lab(*k) for k, v in inside.items() if v["code"] == c]
        L.append(f"  UNKNOWN label {c!r}: {where}")
    blanks_raw = [lab(*k) for k, v in inside.items() if v["code"] == "" and v["raw"] is not None]
    if blanks_raw:
        L.append(f"  cells holding only whitespace (read as fuel): {blanks_raw}")
    got = {"technological channels": sum(cnt.get(c, 0) for c in TECH),
           "CPS rods": sum(cnt.get(c, 0) for c in RODS),
           "DKE(v) channels": cnt.get("X", 0), "start-up chambers (KD)": cnt.get("T", 0)}
    L.append("")
    L.append("AGAINST DOLLEZHAL'S TEXT (p. 57 and the parameter table)")
    for k, e in EXPECT.items():
        L.append(f"  {k:26s} digitised {got[k]:5d}   text {e:5d}   {'OK' if got[k] == e else 'DIFF %+d' % (got[k] - e)}")
    # symmetry: the rod pattern is mirror-symmetric about the grid diagonal (i <-> j)
    spec = {k for k, v in inside.items() if v["code"] in SPECIAL}
    mir = lambda i, j: (j, i)
    L.append("")
    L.append("MIRROR CHECK (the rod pattern is symmetric about the i <-> j diagonal, with exceptions)")
    L.append("  rods whose mirror cell is fuel / absorber / outside -- one side is probably misread:")
    for k in sorted(spec, key=lambda k: (k[1], k[0])):
        m = grid.get(mir(*k))
        if m is None or m["out"] or m["code"] not in SPECIAL:
            tag = "  (confirmed on the scan)" if lab(*k) in CONFIRMED else "  <-- CHECK"
            L.append(f"    {lab(*k)} = {inside[k]['code']:2s}   mirror {lab(*mir(*k))} = "
                     f"{'OUT' if (m is None or m['out']) else (m['code'] or 'fuel')}{tag}")
    L.append("  type differences between mirrored specials (AZ/over-compensation pairs look deliberate):")
    seen = set()
    for k in sorted(spec, key=lambda k: (k[1], k[0])):
        m = mir(*k)
        if m in spec and inside[m]["code"] != inside[k]["code"] and (m, k) not in seen:
            seen.add((k, m))
            L.append(f"    {lab(*k)} = {inside[k]['code']:2s}   {lab(*m)} = {inside[m]['code']}")
    # outline under 90-degree rotation (the outline, unlike the rod types, is 4-fold symmetric)
    N = 47
    rot = lambda i, j: (j, N - i)
    diff = [(lab(*k), lab(*rot(*k))) for k, c in grid.items()
            if c["out"] != grid[rot(*k)]["out"]]
    L.append("")
    L.append("OUTLINE CHECK (in/out under a 90-degree rotation about the grid centre)")
    for a, b in diff:
        ia = [k for k in grid if lab(*k) == a][0]
        L.append(f"    {a} {'OUT' if grid[ia]['out'] else 'in '}  vs  {b}")
    return "\n".join(L) + "\n"


if __name__ == "__main__":
    cols, rows, grid = load()
    open(SVG, "w").write(render(cols, rows, grid))
    rep = checks(cols, rows, grid)
    open(REPORT, "w").write(rep)
    print(rep)
    print("wrote", os.path.relpath(SVG, ROOT), "and", os.path.relpath(REPORT, ROOT))
    if "--png" in sys.argv:
        import cairosvg
        cairosvg.svg2png(url=SVG, write_to=SVG[:-4] + "_preview.png", output_width=1400)
